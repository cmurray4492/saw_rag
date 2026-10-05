from __future__ import annotations

import csv
import hashlib
import logging
import re
import shutil
import threading
import time
from datetime import date, datetime
from datetime import time as time_of_day
from html.parser import HTMLParser
from io import BytesIO, StringIO
from pathlib import Path

from docx import Document as DocxDocument
from docx.table import Table
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader
from watchdog.events import FileSystemEvent, FileSystemEventHandler
from watchdog.observers import Observer

from .chunking import chunk_text
from .config import Config
from .embeddings import Embedder
from .llm import ChatClient
from .stores.base import StoredChunk, VectorStore

log = logging.getLogger(__name__)

TEXT_SUFFIXES = {".txt", ".md", ".markdown"}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
CSV_SUFFIXES = {".csv"}
XLSX_SUFFIXES = {".xlsx", ".xlsm"}
PPTX_SUFFIXES = {".pptx"}
HTML_SUFFIXES = {".html", ".htm"}
DOCUMENT_SUFFIXES = (
    TEXT_SUFFIXES
    | HTML_SUFFIXES
    | PDF_SUFFIXES
    | DOCX_SUFFIXES
    | CSV_SUFFIXES
    | XLSX_SUFFIXES
    | PPTX_SUFFIXES
)
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
SUPPORTED_SUFFIXES = DOCUMENT_SUFFIXES | IMAGE_SUFFIXES

IMAGE_DESCRIBE_PROMPT = (
    "Describe this image in detail for a search index. Include the main "
    "subject, any creatures, people, or objects present, the setting, "
    "mood, colours, and any distinctive visual features. Be factual and "
    "concisse -- a short paragraph is enough. Do not editorialise."
)

def _hash_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()

def _extract_pdf_text(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted and reader.decrypt("") == 0:
        raise ValueError("PDF is password protected")
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)

def _extract_docx_text(data: bytes) -> str:
    document = DocxDocument(BytesIO(data))
    blocks: list[str] = []
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            for row in item.rows:
                # Merged cells repeat the same cell object; keep each once.
                cells = list(dict.fromkeys(row.cells))
                blocks.append(" | ".join(cell.text.strip() for cell in cells))
        else:
            blocks.append(item.text)
    return "\n".join(b for b in blocks if b.strip())

def _rows_to_records(rows: list[list[str]], prefix: str | None = None) -> list[str]:
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return []
    header, body = rows[0], rows[1:]
    if not body:
        line = ", ".join(c.strip() for c in header if c.strip())
        return [f"{prefix}\n{line}" if prefix else line]
    # One paragraph per row so the chunker never splits a row, and each
    # row carries its column names for retrieval.
    def column_name(i: int) -> str:
        if i < len(header) and header[i].strip():
            return header[i].strip()
        return f"column {i + 1}"

    records = []
    for row in body:
        pairs = [
            f"{column_name(i)}: {value.strip()}"
            for i, value in enumerate(row)
            if value.strip()
        ]
        if prefix:
            pairs.insert(0, prefix)
        records.append("\n".join(pairs))
    return records

def _extract_csv_text(data: bytes) -> str:
    # utf-8-sig strips the BOM Excel adds to exported CSVs.
    text = data.decode("utf-8-sig", errors="replace")
    return "\n\n".join(_rows_to_records(list(csv.reader(StringIO(text)))))

def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time_of_day)):
        return value.isoformat()
    return str(value)

def _extract_xlsx_text(data: bytes) -> str:
    # data_only reads cached formula results instead of formula strings.
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        multi_sheet = len(workbook.worksheets) > 1
        records: list[str] = []
        for sheet in workbook.worksheets:
            rows = [
                [_cell_to_str(v) for v in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            prefix = f"sheet: {sheet.title}" if multi_sheet else None
            records.extend(_rows_to_records(rows, prefix))
        return "\n\n".join(records)
    finally:
        workbook.close()

def _pptx_shape_lines(shape) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [line for s in shape.shapes for line in _pptx_shape_lines(s)]
    if shape.has_table:
        return [
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in shape.table.rows
        ]
    if shape.has_text_frame:
        return [shape.text_frame.text]
    return []

def _extract_pptx_text(data: bytes) -> str:
    presentation = Presentation(BytesIO(data))
    slides: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape else ""
        lines = [f"Slide {number}: {title}" if title else f"Slide {number}"]
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            lines.extend(_pptx_shape_lines(shape))
        # has_notes_slide avoids creating an empty notes slide on access.
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                lines.append(f"Notes: {notes.text}")
        # Keep each slide as one paragraph so the chunker keeps it together.
        text = "\n".join(lines).replace("\v", "\n")
        slides.append("\n".join(l for l in text.splitlines() if l.strip()))
    return "\n\n".join(slides)

class _HTMLTextExtractor(HTMLParser):
    _SKIP = {"script", "style", "noscript", "template", "svg", "head"}
    _BLOCK = {
        "p", "div", "section", "article", "main", "aside", "header", "footer",
        "nav", "blockquote", "pre", "table", "ul", "ol", "dl", "form",
        "figure", "figcaption", "h1", "h2", "h3", "h4", "h5", "h6", "hr",
    }
    _LINE = {"br", "li", "tr", "dt", "dd"}
    _CELL = {"td", "th"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.title = ""
        self._skip_depth = 0
        self._in_title = False
        self._in_pre = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag == "title":
            self._in_title = True
        elif tag in self._SKIP:
            self._skip_depth += 1
        elif tag == "pre":
            self._in_pre += 1
            self.parts.append("\n\n")
        elif tag in self._BLOCK:
            self.parts.append("\n\n")
        elif tag in self._LINE:
            self.parts.append("\n")
        elif tag in self._CELL:
            self.parts.append(" | ")

    def handle_startendtag(self, tag: str, attrs) -> None:
        if tag in self._LINE or tag == "hr":
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag == "title":
            self._in_title = False
        elif tag in self._SKIP:
            self._skip_depth = max(0, self._skip_depth - 1)
        elif tag == "pre":
            self._in_pre = max(0, self._in_pre - 1)
            self.parts.append("\n\n")
        elif tag in self._BLOCK:
            self.parts.append("\n\n")

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title += data
        elif self._skip_depth == 0:
            self.parts.append(data if self._in_pre else re.sub(r"\s+", " ", data))

def _extract_html_text(data: bytes) -> str:
    parser = _HTMLTextExtractor()
    parser.feed(data.decode("utf-8-sig", errors="replace"))
    parser.close()
    text = "".join(parser.parts)
    # Trim whitespace and the leading cell separator left by the first <td>.
    lines = [line.strip().removeprefix("| ").strip() for line in text.split("\n")]
    text = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    title = " ".join(parser.title.split())
    return f"{title}\n\n{text}" if title else text

def _is_under(path: Path, root: Path) -> bool:
    return path.resolve().is_relative_to(root.resolve())


class Ingestor:

    def __init__(
            self,
            config: Config,
            store: VectorStore,
            embedder: Embedder,
            chat: ChatClient | None = None,
    ) -> None:
        self.config = config
        self.store = store
        self.embedder = embedder
        self.chat = chat

    def ingest_file(self, path: Path, description: str | None = None) -> None:
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            log.info("Skipping unsupported file: %s", path.name)
            return
        
        try:
            data = path.read_bytes()
        except FileNotFoundError:
            log.warning("File vanished before read: %s", path)
            return
        
        content_hash = _hash_bytes(data)
        source_path = str(path.resolve())

        if self.store.has_document(source_path, content_hash):
            log.info("Unchanged, skipping embed: %s", path.name)
            self._move_to_processed(path)
            return
        
        # Dispatch on suffix
        if suffix in IMAGE_SUFFIXES:
            self._ingest_image(
                path, source_path, content_hash, description=description
            )
        elif suffix in PDF_SUFFIXES:
            self._ingest_text(path, _extract_pdf_text(data), source_path, content_hash)
        elif suffix in DOCX_SUFFIXES:
            self._ingest_text(path, _extract_docx_text(data), source_path, content_hash)
        elif suffix in CSV_SUFFIXES:
            self._ingest_text(path, _extract_csv_text(data), source_path, content_hash)
        elif suffix in XLSX_SUFFIXES:
            self._ingest_text(path, _extract_xlsx_text(data), source_path, content_hash)
        elif suffix in PPTX_SUFFIXES:
            self._ingest_text(path, _extract_pptx_text(data), source_path, content_hash)
        elif suffix in HTML_SUFFIXES:
            self._ingest_text(path, _extract_html_text(data), source_path, content_hash)
        else:
            text = data.decode("utf-8", errors="replace")
            self._ingest_text(path, text, source_path, content_hash)

    def _ingest_text(
            self,
            path: Path,
            text: str,
            source_path: str,
            content_hash: str,
    ) -> None:
        chunks = chunk_text(text, self.config.chunk_size, self.config.chunk_overlap)
        if not chunks:
            log.warning("No content to ingest in %s", path.name)
            self._move_to_processed(path)
            return
        
        log.info("Embedding %d chunks from %s", len(chunks), path.name)
        embeddings = self.embedder.embed([c.text for c in chunks])

        stored = [
            StoredChunk(
                index = c.index,
                text = c.text,
                embedding = emb,
                metadata={"type": "text"},
            )
            for c, emb in zip(chunks, embeddings, strict=True)
        ]

        self.store.upsert_document(
            source_path,
            content_hash,
            stored,
            metadata={"suffix": path.suffix.lower(), "kind":"text"}
        )
        log.info("Ingested %s (%d chunks)", path.name, len(stored))
        self._move_to_processed(path)

    def _ingest_image(
            self,
            path: Path,
            source_path: str,
            content_hash: str,
            description: str | None = None
    ) -> None:
        if description is None:
            if self.chat is None:
                raise RuntimeError(
                    "Image ingestion requires a ChatClient or an explicit "
                    "description; construct the Ingestor with chat=ChatClient(...)"
                )
            log.info(
                "Describing image %s with %s...", path.name, self.config.vision_model
            )
            description = self.chat.describe(
                self.config.vision_model, IMAGE_DESCRIBE_PROMPT, path
            )
        description = description.strip() if description else ""
        if not description:
            log.warning("Empty description for %s; falling back to filename", path.name)
            description = f"Image file: {path.name}"

        chunk_body = f"[Image: {path.name}]\n{description}"
        [embedding] = self.embedder.embed([chunk_body])

        stored = [
            StoredChunk(
                index=0,
                text=chunk_body,
                embedding=embedding,
                metadata={"type":"image"},
            )
        ]
        self.store.upsert_document(
            source_path,
            content_hash,
            stored,
            metadata={"suffix": path.suffix.lower(), "kind": "image"},
        )
        log.info("Ingested image %s", path.name)
        self._move_to_processed(path)

    def _move_to_processed(self, path: Path) -> None:
        processed = self.config.processed_dir
        processed.mkdir(parents=True, exist_ok=True)
        target = processed / path.name
        if target.exists():
            stem, suffix = path.stem, path.suffix
            ts = time.strftime("%Y%m%d-%H%M%S")
            target = processed / f"{stem}.{ts}{suffix}"
        shutil.move(str(path), str(target))
        log.info("Moved -> %s", target.relative_to(self.config.documents_dir.parent))


class _DebouncedHandler(FileSystemEventHandler):

    def __init__(self, ingestor: Ingestor, debounce_seconds: float = 0.75) -> None:
        self._ingestor = ingestor
        self._debounce = debounce_seconds
        self._timers: dict[str, threading.Timer]= {}
        self._lock = threading.Lock()
        self._processed_dir = ingestor.config.processed_dir

    def _schedule(self, raw_path: str) -> None:
        path = Path(raw_path)
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            return
        
        if _is_under(path, self._processed_dir):
            return
        
        with self._lock:
            existing = self._timers.pop(raw_path, None)
            if existing is not None:
                existing.cancel()
            timer = threading.Timer(self._debounce, self._fire, args=(raw_path,))
            timer.daemon = True
            self._timers[raw_path] = timer
            timer.start()

    def _fire(self, raw_path: str) -> None:
        with self._lock:
            self._timers.pop(raw_path, None)

        path = Path(raw_path)
        if not path.exists():
            return
        
        try:
            self._ingestor.ingest_file(path)
        except Exception:
            log.exception("Failed to ingest %s", path)

    def on_created(self, event: FileSystemEvent) -> None:
        if not event.is_directory:
            self._schedule(str(event.src_path))

    def on_moved(self, event: FileSystemEvent) -> None:
        if not event.is_directory:
            self._schedule(str(event.src_path))

    def on_modified(self, event: FileSystemEvent) -> None:
        if not event.is_directory:
            self._schedule(str(event.src_path))


def initial_scan(ingestor: Ingestor) -> None:
    docs = ingestor.config.documents_dir
    processed = ingestor.config.processed_dir
    for path in sorted(docs.iterdir()):
        if path.is_dir():
            continue
        if _is_under(path, processed):
            continue
        try:
            ingestor.ingest_file(path)
        except Exception:
            log.exception("Failed to ingest %s during initial scan", path)


def watch(ingestor: Ingestor) -> None:
    ingestor.config.ensure_dirs()
    docs = ingestor.config.documents_dir

    log.info("Initial scan of %s", docs)
    initial_scan(ingestor)

    handler = _DebouncedHandler(ingestor)
    observer = Observer()
    observer.schedule(handler, str(docs), recursive=False)
    observer.start()
    log.info("Watching %s (Ctrl-C to stop)", docs)
    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        log.info("Stopping watcher...")
    finally:
        observer.stop()
        observer.join()